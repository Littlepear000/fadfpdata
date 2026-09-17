from openpyxl import load_workbook
from openpyxl.styles import Font
from fadfpdata.fm import output_folder

msa_qdrive = output_folder
prefix = 'FMOct2025'
old_date = '20260410'
new_date = '20260414'

# wb1 = load_workbook(fr'{msa_qdrive}\Stat_Tables1-25_{prefix}_{old_date}_copy.xlsx')
# wb2 = load_workbook(fr'{msa_qdrive}\Stat_Tables1-25_{prefix}_{new_date}_copy.xlsx')
wb1 = load_workbook(fr'{msa_qdrive}\Stat_Tables1-25_FMApr2026_20260413_v3.xlsx')
wb2 = load_workbook(fr'{msa_qdrive}\Stat_Tables1-25_FMApr2026_20260414.xlsx')

for sheet_name in wb1.sheetnames:
    print(sheet_name)
    sheet1 = wb1[sheet_name]
    sheet2 = wb2[sheet_name]

    for row in sheet1.iter_rows(min_row=1, max_row=sheet1.max_row, min_col=1, max_col=sheet1.max_column):
        for cell in row:
            cell2 = sheet2[cell.coordinate]

            if cell.value is not None and cell2.value is not None:
                if isinstance(cell.value, (int, float)) and isinstance(cell2.value, (int, float)):
                    if round(cell.value, 1) != round(cell2.value, 1):
                        cell2.font = Font(name='Arial', size=9, color="FF0000")
                else:
                    if cell.value != cell2.value:
                        cell2.font = Font(name='Arial', size=9, color="FF0000")
            elif cell.value is not None or cell2.value is not None:
                cell2.font = Font(name='Arial', size=9, color="FF0000")
            else:
                cell2.font = Font(color="000000")
output_path = fr'{msa_qdrive}\MSA_update_w_red_highlight_{new_date}.xlsx'
wb2.save(output_path)
wb1.close()
wb2.close()

