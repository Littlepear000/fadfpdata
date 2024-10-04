import xlwings as xw


# 打开指定的 Excel 文件
file_path = r'C:\Users\xli7\Desktop\MSA_1004.xlsx'
wb = xw.Book(file_path)

# 遍历所有工作表，并修改名称
for sheet in wb.sheets:
    # 检查工作表名称中是否包含 'Sheet'
    if 'Sheet' in sheet.name:
        # 替换 'Sheet' 为 'STAT'
        new_name = sheet.name.replace('Sheet', 'STAT')
        # 设置工作表的新名称
        sheet.name = new_name


def convert_formulas_to_values(file_path):
    wb = xw.Book(file_path)

    # 遍历所有工作表
    for sheet in wb.sheets:
        # 选择整个工作表
        sheet.api.UsedRange.Copy()  # 复制工作表中使用过的区域
        sheet.api.UsedRange.PasteSpecial(Paste=-4163)  # 粘贴时仅保留值 (-4163 = xlPasteValues)


# 示例用法
convert_formulas_to_values(file_path)


# 定义文件路径
file_path_old = r'C:\Users\xli7\Desktop\MSA_0925.xlsx'
file_path_new = r'C:\Users\xli7\Desktop\MSA_1004_2.xlsx'
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font

wb1 = load_workbook(file_path_old)
wb2 = load_workbook(file_path_new)

# Iterate through each sheet and compare
for sheet_name in wb1.sheetnames:
    # Load sheets
    sheet1 = wb1[sheet_name]
    sheet2 = wb2[sheet_name]

    # Iterate through each cell in the range of sheet1 and sheet2
    for row in sheet1.iter_rows(min_row=1, max_row=sheet1.max_row, min_col=1, max_col=sheet1.max_column):
        for cell in row:
            # Get corresponding cell from sheet2
            cell2 = sheet2[cell.coordinate]

            # Compare cells and change font color if not equal
            if cell.value is not None and cell2.value is not None:
                if isinstance(cell.value, (int, float)) and isinstance(cell2.value, (int, float)):
                    if round(cell.value, 1) != round(cell2.value, 1):
                        # 如果两者近似比较不相等，将 cell2 的字体颜色改为红色
                        cell2.font = Font(name='Arial', size=9, color="FF0000")
            else:
                cell2.font = Font(color="000000")
# Save the modified workbook
output_path = r"C:\Users\xli7\Desktop\MSA_1004_modified.xlsx"
wb2.save(output_path)

